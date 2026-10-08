"""Pruebas de contenido, integridad, cálculos y guía offline de EBA."""
from pathlib import Path
import csv, hashlib, json, re
from urllib.parse import urlparse, unquote
import nbformat
import numpy as np
from scipy import stats
from statsmodels.stats.proportion import proportion_confint
from statsmodels.stats.multitest import multipletests
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright
import pymupdf
from openpyxl import load_workbook
from eba_recursos import *

def main():
    salida=CURSO/'reproducibilidad';cap=salida/'capturas';cap.mkdir(exist_ok=True)
    result={'edicion':'2026-10-08','notebooks':[],'alcance':'Ejecución y contenido local; no aprueba currículo ni automatiza Excel/Power BI Desktop.'}
    for p in sorted((CURSO/'notebooks').glob('*.ipynb')):
        nb=nbformat.read(p,4);nbformat.validate(nb);cells=[c for c in nb.cells if c.cell_type=='code']
        assert cells and all(c.execution_count is not None for c in cells),p.name
        assert not [o for c in cells for o in c.outputs if o.output_type=='error'],p.name
        states=nb.metadata.get('widgets',{}).get('application/vnd.jupyter.widget-state+json',{}).get('state',{});assert states,p.name
        for state in states.values():assert not [o for o in state.get('state',{}).get('outputs',[]) if o.get('output_type')=='error']
        for c in nb.cells:assert not re.search('[\x00-\x08\x0b\x0c\x0e-\x1f]',c.source),p.name
        result['notebooks'].append({'archivo':p.name,'celdas_codigo':len(cells),'resultado':'OK'})
    assert len(result['notebooks'])==8
    paths=[CURSO/'guia_maestra_eba.html',CURSO/'material_propio/EBA_manual_cientifico.html'];soups={p:BeautifulSoup(p.read_text(encoding='utf-8'),'html.parser') for p in paths};links=0
    for p,soup in soups.items():
        ids=[x['id'] for x in soup.find_all(id=True)];assert len(ids)==len(set(ids)),p
        for a in soup.select('a[href]'):
            ref=a['href'];u=urlparse(ref)
            if u.scheme or ref.startswith('//') or not ref:continue
            dest=(p.parent/unquote(u.path)).resolve() if u.path else p
            assert dest.exists(),(p,ref)
            if u.fragment and dest.suffix=='.html':
                s=soups.get(dest) or BeautifulSoup(dest.read_text(encoding='utf-8'),'html.parser');assert s.find(id=unquote(u.fragment)),ref
            links+=1
    with (CURSO/'matriz_cobertura.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.DictReader(f))
    assert {r['criterio_o_tema'] for r in rows if r['nivel']=='criterio'}=={f'{ra}.{ce}' for ra in range(1,4) for ce in range(1,5)}
    assert {r['criterio_o_tema'] for r in rows if r['nivel']=='contenido_minimo'}=={f'M{i:02d}' for i in range(1,18)}
    for r in rows:
        assert soups[paths[0]].find(id=r['html_ancla']) and soups[paths[1]].find(id=r['qmd_ancla']),r
        nb=nbformat.read(CURSO/r['notebook'],4);assert any(f'id="{r["notebook_ancla"]}"' in c.source for c in nb.cells),r
    result['cobertura']={'criterios':12,'minimos':17,'filas':len(rows),'enlaces_html':links}
    # Comprobaciones independientes: fórmulas frente a bibliotecas y casos de frontera.
    assert np.isclose(bayes(.02,.9,.95)['posterior'],180/670)
    assert np.isnan(bayes(0,.9,1)['posterior'])
    for k,n in [(0,20),(1,20),(10,20),(19,20),(20,20),(450,600)]:assert np.allclose(wilson(k,n),proportion_confint(k,n,method='wilson'))
    assert np.allclose(holm([.01,.03,.04]),[.03,.06,.06])
    assert np.allclose(holm([.8,.01,.03,.001]),multipletests([.8,.01,.03,.001],method='holm')[1])
    x=np.array([2.,4.,6.,8.]);ci=ic_media(x);expected=stats.t.interval(.95,len(x)-1,loc=x.mean(),scale=stats.sem(x));assert np.allclose([ci['li'],ci['ls']],expected)
    df=pedidos();grupos=[g.tiempo_min.to_numpy() for _,g in df.groupby('campana')];media=df.tiempo_min.mean();ssb=sum(len(g)*(g.mean()-media)**2 for g in grupos);ssw=sum(((g-g.mean())**2).sum() for g in grupos);f=(ssb/2)/(ssw/(len(df)-3));assert np.isclose(f,anova(df)['F'])
    archivo=CURSO/'material_propio/EBA_tablero_excel.xlsx';wb=load_workbook(archivo,data_only=True)
    assert wb['Pedidos'].max_row==601 and len(wb['Resumen_valores']._charts)==1
    for g,n,tiempo,ticket,pct in wb['Resumen_valores'].iter_rows(min_row=2,values_only=True):
        d=df[df.campana.eq(g)];assert n==len(d) and np.allclose([tiempo,ticket,pct],[d.tiempo_min.mean(),d.ticket_mil.mean(),d.resuelto.mean()*100])
    assert load_workbook(archivo,data_only=False)['Formulas']['C2'].value.startswith('=AVERAGEIF')
    result['calculos']='Bayes, IC t, Wilson, Holm, ANOVA y agregados Excel verificados con casos independientes.'
    with sync_playwright() as pw:
        browser=pw.chromium.launch(channel='chrome',headless=True);page=browser.new_page(viewport={'width':1440,'height':1000});errors=[];requests=[]
        page.on('pageerror',lambda e:errors.append(str(e)))
        def route(r):
            if r.request.url.startswith(('http:','https:')):requests.append(r.request.url);r.abort()
            else:r.continue_()
        page.route('**/*',route);page.goto(paths[0].as_uri(),wait_until='load');assert not errors,errors
        page.screenshot(path=str(cap/'guia_inicio.png'))
        count=0
        for canal in ['Todos','Web','Tienda']:
            for campana in ['Todas','A','B','C']:
                page.select_option('#dash-canal',canal);page.select_option('#dash-campana',campana)
                js=page.evaluate('([c,g])=>EBA.kpis(c,g)',[canal,campana]);d=df
                if canal!='Todos':d=d[d.canal.eq(canal)]
                if campana!='Todas':d=d[d.campana.eq(campana)]
                assert js['n']==len(d);assert np.allclose([js['ticket'],js['tiempo'],js['resueltos']],[d.ticket_mil.mean(),d.tiempo_min.mean(),100*d.resuelto.mean()]);count+=1
        for p,s,e in [(.02,.9,.95),(.001,.5,.999),(.2,.99,.8)]:
            js=page.evaluate('([p,s,e])=>EBA.bayes(p,s,e)',[p,s,e]);assert np.isclose(js['posterior'],bayes(p,s,e)['posterior'])
        for n,p,k in [(20,.8,18),(5,.05,0),(5,.95,50),(50,.95,50),(50,.05,3)]:assert np.isclose(page.evaluate('([n,p,k])=>EBA.binomial(n,p,k)',[n,p,k]),stats.binom.sf(k-1,n,p),atol=1e-12)
        for n in [20,100,600]:
            for c in [.9,.95,.99]:
                js=page.evaluate('([n,c])=>EBA.intervalo(n,c)',[n,c]);py=ic_media(df.tiempo_min.iloc[:n],c);assert np.allclose([js[x] for x in ['media','se','li','ls']],[py[x] for x in ['media','se','li','ls']])
        for e,r in [(.02,.4),(.05,.8),(.15,1)]:
            js=page.evaluate('([e,r])=>EBA.muestra(e,r)',[e,r]);py=tamano_muestra(error=e,respuesta=r);assert js['n']==py['completas'] and js['invitaciones']==py['invitaciones']
        for canal in ['Todos','Web','Tienda']:
            d=df if canal=='Todos' else df[df.canal.eq(canal)];assert np.isclose(page.evaluate('(c)=>EBA.correlacion(c)',canal),stats.pearsonr(d.ticket_mil,d.tiempo_min).statistic)
        y=serie().iloc[:36].to_numpy()
        for alpha in [.05,.3,.95]:
            nivel=[y[0]]
            for v in y[1:]:nivel.append(alpha*v+(1-alpha)*nivel[-1])
            assert np.allclose(page.evaluate('(a)=>EBA.ses(a)',alpha),nivel)
        r=welch(df[df.campana.eq('B')].tiempo_min,df[df.campana.eq('A')].tiempo_min)
        for costo,valor in [(0,.05),(.5,.2),(2,.5)]:
            js=page.evaluate('([c,v])=>EBA.beneficio(c,v)',[costo,valor]);assert np.allclose([js['medio'],js['li'],js['ls']],[-r['diferencia']*valor-costo,-r['ls']*valor-costo,-r['li']*valor-costo])
        changes=0
        for control in page.locator('.controles input,.controles select').all():
            id=control.get_attribute('id');tag=control.evaluate('(x)=>x.tagName')
            if tag=='SELECT':
                options=control.locator('option').evaluate_all('(xs)=>xs.map(x=>x.value)')
                for v in options:control.select_option(v);changes+=1
            else:
                for key in ['min','max']:
                    control.fill(control.get_attribute(key));control.dispatch_event('input');changes+=1
        assert not errors,errors
        for i in range(1,9):
            q=page.locator(f'#quiz-{i}');q.locator('button').click();assert 'Selecciona' in q.locator('p').inner_text()
            q.locator('input[value="0"]').check();q.locator('button').click();assert 'Revisa.' in q.locator('p').inner_text()
            q.locator('input[value="1"]').check();q.locator('button').click();assert 'Correcto.' in q.locator('p').inner_text()
        page.fill('#buscar','zzzzsinresultado');assert page.locator('#sin-resultados').is_visible()
        page.fill('#buscar','Bayes');assert page.locator('#sec-bayes').is_visible()
        page.locator('aside nav a[href="#sec-anova"]').click();assert page.input_value('#buscar')=='' and page.locator('#sec-problema').is_visible()
        page.locator('#lab-bayes').scroll_into_view_if_needed();page.screenshot(path=str(cap/'guia_bayes.png'))
        page.set_viewport_size({'width':390,'height':844});page.goto(paths[0].as_uri());assert page.evaluate('document.documentElement.scrollWidth<=innerWidth'),page.evaluate('document.documentElement.scrollWidth')
        page.screenshot(path=str(cap/'guia_movil.png'))
        assert not errors and not requests,(errors,requests)
        result['guia']={'sin_red':True,'errores_js':errors,'filtros_conciliados':count,'cambios_controles':changes,'autoevaluaciones':8,'movil_sin_desborde':True}
        browser.close()
    doc=pymupdf.open(CURSO/'material_propio/EBA_manual_cientifico.pdf');text='\n'.join(p.get_text() for p in doc)
    assert len(doc)>=15 and len(text)>35000,(len(doc),len(text))
    for term in ['Bayes','ANOVA','ARIMA','Wilson','Holm','Power BI','regresión','muestreo']:assert term in text,term
    assert len(soups[paths[1]].find_all('math'))>=20
    for i in [0,4,len(doc)-2]:doc[i].get_pixmap(matrix=pymupdf.Matrix(1.2,1.2)).save(cap/f'manual_p{i+1:02d}.png')
    result['manual']={'paginas':len(doc),'caracteres':len(text),'formulas_mathml':len(soups[paths[1]].find_all('math'))}
    (salida/'verificacion_edicion.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
