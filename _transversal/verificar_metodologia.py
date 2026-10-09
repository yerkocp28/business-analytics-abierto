"""Regresiones de los defectos detectados en la auditoría del 09-10-2026.

Ejecuta funciones reales y ejercicios publicados con contraejemplos pequeños.
No sustituye la ejecución completa de notebooks ni una revisión académica externa.
"""
from pathlib import Path
import ast
from collections import Counter
import contextlib
import io
import sys
import unittest

import nbformat
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.linear_model import Ridge
from sklearn.model_selection import GridSearchCV, KFold
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
for carpeta in ['03_Modelos_Predictivos_Negocios', '04_Analitica_Estrategica_Datos']:
    sys.path.insert(0, str(RAIZ / carpeta / 'reproducibilidad'))
import mpn_recursos as mr
import aed_recursos as ar
from _transversal.evaluacion import pregunta
from _transversal.rutas import CURSOS, modulos


def notebook(curso, numero):
    return nbformat.read(next((RAIZ / CURSOS[curso]['carpeta'] / 'notebooks').glob(f'{curso}_{numero:02d}_*.ipynb')), 4)


class Metodologia(unittest.TestCase):
    def test_psi_binaria_con_cambio_de_prevalencia(self):
        esperado = [0] * 95 + [1] * 5
        actual = [0] * 20 + [1] * 80
        manual = (.8 - .05) * np.log(.8 / .05) + (.2 - .95) * np.log(.2 / .95)
        self.assertAlmostEqual(mr.psi(esperado, actual), manual)

    def test_psi_constantes_categorias_nuevas_y_faltantes(self):
        self.assertEqual(mr.psi([1] * 10, [1] * 10), 0)
        self.assertGreater(mr.psi([1] * 10, [2] * 10), 20)
        self.assertGreater(mr.psi(['A', 'B'], ['A', 'C']), 0)
        self.assertGreater(mr.psi([1, 1, None], [1, None, None]), 0)
        self.assertEqual(mr.psi([None, None], [None, None]), 0)
        with self.assertRaises(ValueError):
            mr.psi([], [1])
        with self.assertRaises(ValueError):
            mr.psi([1, 1], [1, 2], tipo='continuo')

    def test_psi_continuo_incluye_colas_y_faltantes(self):
        x = np.arange(100, dtype=float)
        self.assertEqual(mr.psi(x, x, tipo='continuo'), 0)
        self.assertGreater(mr.psi(x, x + 1000, tipo='continuo'), 0)
        self.assertGreater(mr.psi(x, np.r_[x[:50], [np.nan] * 50], tipo='continuo'), 0)

    def test_calibracion_independiente(self):
        x = pd.DataFrame({'x': np.arange(480)}, index=np.arange(1000, 1480))
        y = x.x * 2
        partes = mr.particion_con_calibracion(x, y)
        conjuntos = [set(p.index) for p in partes[:4]]
        self.assertEqual([len(s) for s in conjuntos], [288, 48, 48, 96])
        self.assertEqual(set.union(*conjuntos), set(x.index))
        for i in range(4):
            self.assertTrue(partes[i].index.equals(partes[i + 4].index))
            for j in range(i):
                self.assertFalse(conjuntos[i] & conjuntos[j])

    def test_cuantil_conformal_finito(self):
        self.assertEqual(mr.cuantil_conformal(np.arange(1, 10), .8), 8)
        self.assertEqual(mr.cuantil_conformal([1, 2, 3], .99), np.inf)
        for residuos in [[], [-1, 1], [1, np.nan]]:
            with self.assertRaises(ValueError):
                mr.cuantil_conformal(residuos)

    def test_escalador_se_ajusta_en_cada_pliegue_del_notebook(self):
        celda = next(c.source for c in notebook('MPN', 3).cells
                     if c.cell_type == 'code' and 'def seleccionar_regularizacion' in c.source)
        funcion = next(n for n in ast.parse(celda).body if isinstance(n, ast.FunctionDef))
        tamanos = []

        class EscaladorObservado(StandardScaler):
            def fit(self, X, y=None, **kwargs):
                tamanos.append(len(X))
                return super().fit(X, y, **kwargs)

        rng = np.random.default_rng(0)
        x = rng.normal(size=(100, 3))
        entorno = dict(Pipeline=Pipeline, StandardScaler=EscaladorObservado, GridSearchCV=GridSearchCV,
                       cv_interna=KFold(5, shuffle=True, random_state=0), Ae=x, ye=x[:, 0] + rng.normal(size=100))
        exec(compile(ast.Module(body=[funcion], type_ignores=[]), 'seleccionar_regularizacion', 'exec'), entorno)
        entorno['seleccionar_regularizacion'](Ridge(), [.1, 1])
        self.assertEqual(tamanos[:-1], [80] * 10)
        self.assertEqual(tamanos[-1], 100)  # refit final, una vez seleccionada alpha

    def test_intervalos_no_usan_resultados_futuros(self):
        celda = next(c.source for c in notebook('MPN', 7).cells
                     if c.cell_type == 'code' and 'primer_origen_verificacion' in c.source)

        def origen_movil(h):
            origenes = np.arange(36, 73 - h)
            return pd.DataFrame({'origen': origenes, 'método': 'Tendencia + estacionalidad',
                                 'error_ultimo_paso': origenes.astype(float)})

        entorno = dict(np=np, pd=pd, origen_movil=origen_movil, display=lambda x: None,
                       serie=pd.Series(np.arange(72), index=pd.date_range('2020-01-01', periods=72, freq='MS')))
        exec(compile(celda, 'intervalos_publicados', 'exec'), entorno)
        for h, fila in entorno['intervalos'].iterrows():
            self.assertLess(pd.Timestamp(fila['último resultado de calibración']), pd.Timestamp(fila['primera emisión']))
            r = origen_movil(h)
            corte = r.origen.iloc[len(r) // 2]
            esperados = int(((r.origen + h - 1) < corte).sum())
            self.assertEqual(fila['errores de calibración disponibles'], esperados)
            if h > 1:
                self.assertLess(esperados, len(r) // 2)

    def test_boldonet_expone_todas_las_incidencias(self):
        df, a = ar.boldonet()
        for campo, esperado in {'actualizaciones_anteriores_ingreso': 21, 'actualizaciones_invalidas': 1,
                               'satisfacciones_fuera_rango': 4, 'permanencias_negativas': 2,
                               'elegibles_kpi': 121, 'elegibles_kpi_estricto': 121}.items():
            self.assertEqual(a[campo], esperado)
        self.assertTrue(df.loc[df.elegible_kpi_estricto, 'maduro'].all())
        self.assertFalse(df.loc[df.elegible_kpi_estricto, 'actualizacion_antes_ingreso'].any())

    def test_todas_las_alternativas_de_los_28_quizzes(self):
        posiciones = Counter()
        cantidad = 0
        for codigo in CURSOS:
            for numero in range(1, len(CURSOS[codigo]['modulos']) + 1):
                nb = notebook(codigo, numero)
                llamadas = []
                for c in nb.cells:
                    if c.cell_type == 'code':
                        arbol = ast.parse(c.source.replace('%matplotlib inline', '# magic'))
                        llamadas.extend(n for n in ast.walk(arbol) if isinstance(n, ast.Call)
                                        and (getattr(n.func, 'attr', '') == 'pregunta'
                                             or getattr(n.func, 'id', '') == 'pregunta'))
                self.assertEqual(len(llamadas), 1, (codigo, numero))
                args = [ast.literal_eval(a) for a in llamadas[0].args]
                if codigo == 'EBA':
                    if len(args) == 4:
                        args.append(next(ast.literal_eval(k.value) for k in llamadas[0].keywords if k.arg == 'tercera'))
                    texto, incorrecta, correcta, explicacion, tercera = args
                    args = [texto, [incorrecta, correcta, tercera], 1, explicacion]
                with contextlib.redirect_stdout(io.StringIO()):
                    caja = pregunta(*args, mostrar=False)
                    radio, boton, salida = caja.children[1:]
                    valores = [v for _, v in radio.options]
                    self.assertGreaterEqual(len(valores), 3)
                    posiciones[valores.index(caja._ba_clave)] += 1
                    for valor in valores:
                        radio.value = valor
                        boton.click()
                        self.assertIn('Correcto' if valor == caja._ba_clave else 'Revisa', salida.value)
                caja.close()
                cantidad += 1
        self.assertEqual(cantidad, 28)
        self.assertGreater(len(posiciones), 1)
        print('Posiciones de respuestas correctas (base 0):', dict(posiciones))

    def test_libros_conservan_formulas_y_cache_numerico(self):
        fba = RAIZ / CURSOS['FBA']['carpeta'] / 'material_propio/FBA_herramientas.xlsx'
        form, val = load_workbook(fba, data_only=False), load_workbook(fba, data_only=True)
        for fila in range(2, val['Resumen'].max_row + 1):
            for col in ['B', 'C', 'D']:
                self.assertEqual(form['Resumen'][f'{col}{fila}'].data_type, 'f')
                self.assertAlmostEqual(val['Resumen'][f'{col}{fila}'].value, val['Control_valores'][f'{col}{fila}'].value)
        eba = RAIZ / CURSOS['EBA']['carpeta'] / 'material_propio/EBA_tablero_excel.xlsx'
        form, val = load_workbook(eba, data_only=False), load_workbook(eba, data_only=True)
        df = pd.read_csv(RAIZ / CURSOS['EBA']['carpeta'] / 'datos/comerciosur_pedidos.csv')
        for fila, campana in enumerate(['A', 'B', 'C'], 2):
            grupo = df.loc[df.campana.eq(campana)]
            self.assertEqual(form['Formulas'][f'B{fila}'].data_type, 'f')
            self.assertEqual(val['Formulas'][f'B{fila}'].value, len(grupo))
            self.assertAlmostEqual(val['Formulas'][f'C{fila}'].value, grupo.tiempo_min.mean())


if __name__ == '__main__':
    unittest.main(verbosity=2)
