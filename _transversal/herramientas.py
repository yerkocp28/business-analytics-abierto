"""Libros propios y valores precalculados; no pretende ejecutar Excel Desktop."""
from pathlib import Path
from datetime import datetime
import io
import zipfile
import xml.etree.ElementTree as ET
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parents[1]
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def guardar_cache_formulas(archivo, valores):
    """Inserta resultados calculados en Python para lectores que no evalúan fórmulas.

    Las fórmulas se conservan; al abrir en Excel se solicita recalcular. El caché
    facilita lectura, no certifica haber ejecutado Excel ni el motor de Power BI.
    """
    archivo = Path(archivo)
    libro = load_workbook(archivo, data_only=False)
    hojas = {nombre: i + 1 for i, nombre in enumerate(libro.sheetnames)}
    with zipfile.ZipFile(archivo) as z:
        partes = {nombre: z.read(nombre) for nombre in z.namelist()}
    for nombre, celdas in valores.items():
        ruta = f'xl/worksheets/sheet{hojas[nombre]}.xml'
        raiz = ET.fromstring(partes[ruta])
        for referencia, valor in celdas.items():
            celda = raiz.find(f'.//{{{NS}}}c[@r="{referencia}"]')
            if celda is None or celda.find(f'{{{NS}}}f') is None:
                raise ValueError(f'No existe fórmula en {nombre}!{referencia}')
            cache = celda.find(f'{{{NS}}}v')
            if cache is None:
                cache = ET.SubElement(celda, f'{{{NS}}}v')
            cache.text = repr(float(valor))
        partes[ruta] = ET.tostring(raiz, encoding='utf-8', xml_declaration=True)
    # Metadatos y ZIP estables para reconstrucciones comparables.
    core = ET.fromstring(partes['docProps/core.xml'])
    for tag in ['created', 'modified']:
        nodo = core.find('{http://purl.org/dc/terms/}' + tag)
        if nodo is not None:
            nodo.text = '2026-10-09T00:00:00Z'
    partes['docProps/core.xml'] = ET.tostring(core, encoding='utf-8', xml_declaration=True)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as z:
        for nombre in sorted(partes):
            info = zipfile.ZipInfo(nombre, (2026, 10, 9, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, partes[nombre])
    archivo.write_bytes(buffer.getvalue())


def crear_libro_fba():
    datos = pd.read_csv(RAIZ / '_transversal/datos/quillaymarket/originales/quillaymarket_ventas_detalle.csv', parse_dates=['Fecha'])
    libro = Workbook()
    libro.properties.creator = 'Colección abierta de Business Analytics'
    libro.properties.lastModifiedBy = 'Colección abierta de Business Analytics'
    hoja = libro.active
    hoja.title = 'Datos'
    hoja.append(list(datos.columns))
    for fila in datos.itertuples(index=False, name=None):
        hoja.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else v for v in fila])
    tabla = Table(displayName='Ventas', ref=f'A1:K{len(datos)+1}')
    tabla.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    hoja.add_table(tabla)
    resumen = libro.create_sheet('Resumen')
    resumen.append(['Canal', 'Ingreso CLP', 'Transacciones', 'Ticket CLP'])
    control = libro.create_sheet('Control_valores')
    control.append(['Canal', 'Ingreso CLP', 'Transacciones', 'Ticket CLP'])
    cache = {'Resumen': {}}
    fin = len(datos) + 1
    for i, (canal, grupo) in enumerate(datos.groupby('Canal'), 2):
        resumen.append([canal, f'=SUMIF(Datos!H2:H{fin},A{i},Datos!K2:K{fin})',
                        f'=COUNTIF(Datos!H2:H{fin},A{i})', f'=IFERROR(B{i}/C{i},0)'])
        ingreso, n = float(grupo.Monto.sum()), len(grupo)
        control.append([canal, ingreso, n, ingreso/n])
        cache['Resumen'].update({f'B{i}': ingreso, f'C{i}': n, f'D{i}': ingreso/n})
    total = resumen.max_row + 1
    resumen.append(['Total', f'=SUM(B2:B{total-1})', f'=SUM(C2:C{total-1})', f'=B{total}/C{total}'])
    cache['Resumen'].update({f'B{total}': float(datos.Monto.sum()), f'C{total}': len(datos),
                            f'D{total}': float(datos.Monto.mean())})
    control.append(['Total', float(datos.Monto.sum()), len(datos), float(datos.Monto.mean())])
    instrucciones = libro.create_sheet('Instrucciones')
    for fila in [
        ['Paso', 'Trabajo y comprobación'],
        [1, 'Datos: 180 ventas sintéticas, enero-junio de 2025; una fila es una transacción, no un cliente.'],
        [2, 'Resumen conserva fórmulas; Control_valores contiene el contraste independiente en Python.'],
        [3, 'Recalcular en Excel y comparar ambas hojas. Los valores cacheados provienen de Python.'],
        [4, 'Los filtros visuales de Datos no cambian automáticamente SUMIF/COUNTIF: definir el mismo universo en la fórmula.'],
        [5, 'Crear una tabla dinámica filtrada a abril-junio y conciliar ingreso y transacciones con pandas.'],
        [6, 'Ticket agregado: ingreso total / transacciones; no promedio simple de tickets por canal.'],
        [7, 'Actividad: duplicar una clave del catálogo en una copia y justificar el control many_to_one.'],
        [8, 'Licencias: material/datos CC BY-NC-SA 4.0; generador MIT. No contiene datos de empresas reales.']]:
        instrucciones.append(fila)
    for ws in libro:
        ws.freeze_panes = 'A2'
        for c in ws[1]:
            c.font = Font(color='FFFFFF', bold=True)
            c.fill = PatternFill('solid', fgColor='175159')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
    instrucciones.column_dimensions['B'].width = 110
    for fila in hoja.iter_rows(min_row=2, min_col=2, max_col=2):
        fila[0].number_format = 'yyyy-mm-dd'
    destino = RAIZ / '01_Fundamentos_Business_Analytics/material_propio/FBA_herramientas.xlsx'
    libro.save(destino)
    guardar_cache_formulas(destino, cache)
    return destino


def crear_controles_bi():
    """Contrastes independientes de las medidas DAX, por total y por filtro."""
    import sys
    sys.path.insert(0, str(RAIZ / '04_Analitica_Estrategica_Datos/reproducibilidad'))
    from aed_recursos import casapeumo, indicadores

    pedidos = pd.read_csv(RAIZ / '02_Estadistica_Business_Analytics/datos/comerciosur_pedidos.csv')
    ventas, _, _ = casapeumo()
    for carpeta, df, columna in [
        ('02_Estadistica_Business_Analytics', pedidos, 'canal'),
        ('04_Analitica_Estrategica_Datos', ventas, 'Canal'),
    ]:
        filas = []
        grupos = [('TODOS', df)] + list(df.groupby(columna))
        for filtro, grupo in grupos:
            if carpeta.startswith('02'):
                medidas = {'Pedidos observados': len(grupo), 'Tiempo medio min': grupo.tiempo_min.mean(),
                           'Ticket medio mil UM': grupo.ticket_mil.mean(), 'Resolución': grupo.resuelto.mean()}
            else:
                k = indicadores(grupo)
                medidas = dict(zip(['Ventas únicas', 'Ingresos UM', 'Margen UM', 'Margen relativo',
                                    'Ticket UM', 'Entrega media días'],
                                   [k['ventas'], k['ingreso'], k['margen'], k['margen_pct'], k['ticket'], k['entrega_media']]))
            filas.extend({'columna_filtro': columna, 'valor_filtro': filtro, 'medida': m, 'valor_python': v}
                         for m, v in medidas.items())
        destino = RAIZ / carpeta / 'material_propio/herramientas_bi/controles.csv'
        pd.DataFrame(filas).to_csv(destino, index=False, lineterminator='\n')


if __name__ == '__main__':
    print('Libro propio:', crear_libro_fba().name)
    crear_controles_bi()
    print('Controles de medidas BI preparados con Python.')
